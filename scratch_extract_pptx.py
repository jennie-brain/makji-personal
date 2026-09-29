import sys
import glob
from pptx import Presentation

import os
all_pptx = [p for p in glob.glob(r"C:\Users\Administrator\Downloads\*.pptx") if not os.path.basename(p).startswith("~$")]
candidates = [p for p in all_pptx if "MAKJI_STOCK" in p and "\uc784\uc6d0" in p]
path = candidates[0]
print("Resolved path repr:", repr(path))
out_path = r"C:\Users\Administrator\workspace\makji\scratch_exec_report_utf8.md"

prs = Presentation(path)

def walk(shapes, f, depth=0):
    for shape in shapes:
        try:
            if shape.shape_type == 6:  # GROUP
                walk(shape.shapes, f, depth+1)
                continue
        except Exception:
            pass
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            t = shape.text_frame.text
            if t and t.strip():
                f.write(t.replace("\x0b", "\n") + "\n")
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [cell.text for cell in row.cells]
                f.write(" | ".join(cells) + "\n")
        if getattr(shape, "has_chart", False) and shape.has_chart:
            chart = shape.chart
            try:
                cats = list(chart.plots[0].categories)
                f.write("[CHART categories] " + ", ".join(str(c) for c in cats) + "\n")
            except Exception:
                pass
            for series in chart.series:
                try:
                    f.write(f"[CHART series {series.name}] " + ", ".join(str(v) for v in series.values) + "\n")
                except Exception:
                    pass

with open(out_path, "w", encoding="utf-8") as f:
    for i, slide in enumerate(prs.slides, start=1):
        f.write(f"\n<!-- Slide number: {i} -->\n")
        walk(slide.shapes, f)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            f.write("[NOTES] " + slide.notes_slide.notes_text_frame.text + "\n")

print("done", out_path)
