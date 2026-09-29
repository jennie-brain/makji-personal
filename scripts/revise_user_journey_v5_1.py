"""Create a non-destructive deck variant with a refined purchase-journey slide."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "04_presentations" / "막지스톡_기업대표_서비스기획_발표v5.0.pptx"
OUTPUT = ROOT / "docs" / "04_presentations" / "막지스톡_기업대표_서비스기획_발표v5.1_구매여정개선.pptx"
JOURNEY_INDEX = 6  # slide 7

FONT = "Noto Sans KR"
INK = RGBColor(35, 31, 27)
INK2 = RGBColor(93, 83, 75)
MUTE = RGBColor(151, 143, 134)
LINE = RGBColor(231, 225, 218)
WHITE = RGBColor(255, 255, 255)
CREAM = RGBColor(251, 248, 243)
BREAD = RGBColor(177, 119, 70)
BREAD_DARK = RGBColor(132, 85, 48)
BREAD_LIGHT = RGBColor(246, 232, 217)
SAGE = RGBColor(101, 151, 126)
SAGE_LIGHT = RGBColor(232, 243, 237)


def clear(slide):
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def shape(slide, kind, x, y, w, h, fill=None, line_color=None, width=0.75):
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        item.fill.background()
    else:
        item.fill.solid(); item.fill.fore_color.rgb = fill
    if line_color is None:
        item.line.fill.background()
    else:
        item.line.color.rgb = line_color; item.line.width = Pt(width)
    return item


def text(slide, x, y, w, h, value, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    item = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = item.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.alignment = align
        p.font.name = FONT; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
    return item


def line(slide, x1, y1, x2, y2, color=LINE, width=1.0):
    item = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    item.line.color.rgb = color; item.line.width = Pt(width)
    return item


def arrow(slide, x, y, w, h, color=BREAD):
    return shape(slide, MSO_SHAPE.RIGHT_ARROW, x, y, w, h, color)


def rounded(slide, x, y, w, h, fill=WHITE, line_color=None):
    item = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line_color)
    item.adjustments[0] = 0.10
    return item


def circle(slide, x, y, d, fill):
    return shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill)


prs = Presentation(SOURCE)
slide = prs.slides[JOURNEY_INDEX]
clear(slide)

# Header: one concise, non-wrapping statement.
text(slide, 0.9, 0.56, 8.0, 0.25, "사용자 여정 · 구매 의사별", 11, True, BREAD_DARK)
text(slide, 0.9, 0.91, 11.55, 0.48, "상황이 달라도 고객은 결국 자사몰에서 구매합니다", 25, True)

# Source-state card: deliberately empty center for a future application screenshot.
rounded(slide, 0.9, 1.84, 2.28, 4.13, CREAM, LINE)
text(slide, 1.16, 2.12, 1.75, 0.28, "1 · 오늘의 가격 확인", 11, True, BREAD_DARK, PP_ALIGN.CENTER)
rounded(slide, 1.2, 2.62, 1.68, 2.36, WHITE, RGBColor(238, 221, 203))
text(slide, 1.38, 3.32, 1.32, 0.28, "실제 앱 화면\n목업 영역", 11, True, MUTE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(slide, 1.16, 5.22, 1.75, 0.34, "오늘 가격을 보고\n내 상황에 맞게 선택", 10.5, False, INK2, PP_ALIGN.CENTER)

arrow(slide, 3.35, 3.73, 0.47, 0.28, BREAD)

# Decision lanes: three purchase intents. The end state is a shared mall purchase.
text(slide, 4.05, 1.88, 5.2, 0.28, "2 · 구매 의사에 따라 한 가지 길을 고릅니다", 12, True, INK2)

lanes = [
    (2.30, SAGE_LIGHT, SAGE, "지금 살래요", "오늘 가격으로\n바로 구매", "오늘가로\n즉시 결제"),
    (3.47, BREAD_LIGHT, BREAD, "가격은 좋은데\n결제는 이따가", "오전가 1회 잠금", "16:00부터 01:59까지\n잠근가 또는 더 싼가"),
    (4.64, CREAM, BREAD_DARK, "더 낮아지면\n살래요", "내일 가격 예측", "적중 시 쿠폰으로\n다음 방문·구매"),
]
for idx, (y, fill, accent, intent, action, outcome) in enumerate(lanes, 1):
    rounded(slide, 4.05, y, 7.15, 0.91, fill, None)
    circle(slide, 4.27, y + 0.27, 0.36, accent)
    text(slide, 4.27, y + 0.27, 0.36, 0.36, str(idx), 10, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(slide, 4.82, y + 0.17, 1.46, 0.55, intent, 12.5, True, INK, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    arrow(slide, 6.44, y + 0.32, 0.39, 0.22, accent)
    rounded(slide, 6.99, y + 0.14, 1.56, 0.62, WHITE, None)
    text(slide, 7.07, y + 0.14, 1.4, 0.62, action, 10.5, True, accent, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    arrow(slide, 8.73, y + 0.32, 0.39, 0.22, accent)
    rounded(slide, 9.28, y + 0.14, 1.67, 0.62, WHITE, None)
    text(slide, 9.35, y + 0.14, 1.53, 0.62, outcome, 9.3, True, INK2, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# Shared destination: intentionally open box for a future mall/app mock-up.
text(slide, 11.52, 1.88, 0.9, 0.28, "3 · 자사몰", 11, True, BREAD_DARK, PP_ALIGN.CENTER)
rounded(slide, 11.42, 2.30, 1.0, 3.25, WHITE, RGBColor(238, 221, 203))
text(slide, 11.57, 3.32, 0.70, 0.44, "실제\n구매 화면", 10, True, MUTE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
rounded(slide, 11.54, 4.66, 0.76, 0.35, BREAD, None)
text(slide, 11.54, 4.66, 0.76, 0.35, "결제", 10, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
for y, c in ((2.72, SAGE), (3.89, BREAD), (5.06, BREAD_DARK)):
    arrow(slide, 11.02, y, 0.30, 0.20, c)

# Evidence band: concise internal behavior evidence, no extrapolation.
rounded(slide, 0.9, 6.26, 11.53, 0.52, CREAM, None)
circle(slide, 1.16, 6.41, 0.17, BREAD)
text(slide, 1.48, 6.34, 10.6, 0.33, "5월 자사몰: 오전 8~10시 로그인 25.6% · 결제금액 8.0%  →  저녁 20~24시 로그인 20.0% · 결제금액 28.4%", 10.2, True, INK2, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)

# So-what and provenance.
text(slide, 3.35, 6.93, 9.08, 0.32, "가격 잠금은 아침의 관심을 저녁의 결제로 잇고, 예측은 다음 날 다시 찾게 합니다.", 13, True, BREAD_DARK, PP_ALIGN.RIGHT)
text(slide, 0.9, 7.18, 10.4, 0.18, "출처: 막지 자사몰 5월 시간별 분석(로그인 90회 · 결제금액 기준, 단일상품 막지 ZERO 카스테라) · 가격 잠금 운영 규칙: 브레드마켓 PRD", 7.5, False, MUTE)
text(slide, 11.1, 7.18, 1.33, 0.18, "07 / 18", 8.5, False, MUTE, PP_ALIGN.RIGHT)

slide.notes_slide.notes_text_frame.text = (
    "[핵심] 상황이 달라도 고객은 결국 자사몰에서 구매합니다.\n"
    "[근거] 5월 자사몰에서는 오전 8~10시 로그인 비중이 25.6%지만 결제금액 비중은 8.0%이고, 저녁 20~24시는 로그인 20.0%에 결제금액 28.4%입니다. 즉시 구매형은 오늘 가격으로 바로 사고, 가격 확보형은 오전가를 하루 한 번 잠근 뒤 16:00~01:59에 잠근가 또는 더 싼가로 결제합니다. 관망형은 내일 가격을 예측해 적중 시 쿠폰을 받고 다시 방문합니다.\n"
    "[연결] 가격 잠금은 구매를 미루게 하는 장치가 아니라, 아침의 관심을 저녁의 결제로 연결하는 장치입니다."
)

prs.save(OUTPUT)
print(OUTPUT)
