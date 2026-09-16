from pathlib import Path
import re

from lxml import html
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "04_presentations" / "막지스톡_기업대표_서비스기획_발표v0.1.html"
OUTPUT = SOURCE.with_suffix(".pptx")

SW, SH = 13.333333, 7.5
SKY = RGBColor(66, 165, 245)
SKY2 = RGBColor(11, 121, 208)
SKY3 = RGBColor(223, 242, 255)
SKY4 = RGBColor(244, 251, 255)
INK = RGBColor(16, 35, 59)
MUTED = RGBColor(95, 113, 134)
LINE = RGBColor(216, 231, 243)
WHITE = RGBColor(255, 255, 255)
NAVY = RGBColor(23, 57, 95)
GOLD = RGBColor(240, 169, 58)
FONT = "맑은 고딕"


def clean_text(node):
    if node is None:
        return ""
    text = "\n".join(t.strip() for t in node.itertext() if t.strip())
    return re.sub(r"[ \t]+", " ", text).strip()


def add_rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.8)
    return shape


def set_shape_text(shape, text, size=14, color=INK, bold=False, margin=0.12,
                   valign=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return shape


def add_text(slide, text, x, y, w, h, size=14, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    return set_shape_text(box, text, size, color, bold, 0, valign, align)


def add_rich_block(slide, title, body, x, y, w, h, accent=False, strong=False):
    fill = SKY2 if strong else (SKY4 if accent else WHITE)
    line = SKY2 if strong else LINE
    shape = add_rect(slide, x, y, w, h, fill, line)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(0.19)
    tf.margin_top = tf.margin_bottom = Inches(0.15)
    p = tf.paragraphs[0]
    p.text = title or " "
    p.font.name = FONT
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = WHITE if strong else NAVY
    if body:
        p2 = tf.add_paragraph()
        p2.text = body
        p2.space_before = Pt(6)
        p2.font.name = FONT
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = RGBColor(234, 247, 255) if strong else MUTED
        p2.line_spacing = 1.12
    return shape


def add_header(slide, section, eyebrow, title, page):
    add_rect(slide, 0, 0, 0.10, SH, SKY2, SKY2, False)
    add_text(slide, "MAKJI STOCK", 0.78, 0.45, 2.3, 0.25, 10, SKY2, True)
    add_text(slide, section, 9.1, 0.45, 3.45, 0.25, 9, MUTED, False, PP_ALIGN.RIGHT)
    add_rect(slide, 0.78, 0.91, 0.25, 0.03, SKY, SKY, False)
    add_text(slide, eyebrow, 1.11, 0.80, 4.8, 0.25, 10, SKY2, True)
    add_text(slide, title, 0.78, 1.12, 11.75, 1.02, 25, INK, True)
    add_text(slide, f"{page:02d}", 12.08, 7.06, 0.5, 0.18, 9, MUTED, False, PP_ALIGN.RIGHT)


def node_blocks(section):
    candidates = section.xpath(
        ".//*[contains(concat(' ',normalize-space(@class),' '),' card ') or "
        "contains(concat(' ',normalize-space(@class),' '),' node ') or "
        "contains(concat(' ',normalize-space(@class),' '),' pipe ') or "
        "contains(concat(' ',normalize-space(@class),' '),' decision ') or "
        "contains(concat(' ',normalize-space(@class),' '),' risk ')]"
    )
    out = []
    for n in candidates:
        if any(a in candidates for a in n.iterancestors()):
            continue
        heads = n.xpath(".//h3|.//b[1]")
        title = clean_text(heads[0]) if heads else ""
        body_parts = []
        for e in n.xpath(".//p|.//li|.//span[contains(@class,'num')]"):
            t = clean_text(e)
            if t and t != title and t not in body_parts:
                body_parts.append(t)
        body = "\n".join(("• " + t if e_tag == 'li' else t) for t, e_tag in [
            (clean_text(e), e.tag.lower()) for e in n.xpath(".//p|.//li|.//span[contains(@class,'num')]")
            if clean_text(e) and clean_text(e) != title
        ])
        cls = n.get("class", "")
        out.append((title, body, "blue" in cls, "strong" in cls))
    return out


def cover_slide(prs, section, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = SKY4
    add_rect(slide, 0, 0, SW, 0.08, SKY2, SKY2, False)
    eyebrow = clean_text((section.xpath(".//*[contains(@class,'eyebrow')]") or [None])[0])
    h1 = clean_text((section.xpath(".//h1") or [None])[0])
    h2 = clean_text((section.xpath(".//h2") or [None])[0])
    lead = clean_text((section.xpath(".//*[contains(@class,'lead')]") or [None])[0])
    add_text(slide, eyebrow, 0.80, 1.12, 5.8, 0.32, 11, SKY2, True)
    add_text(slide, h1 or h2, 0.80, 1.58, 7.3, 1.45, 38 if h1 else 31, INK, True)
    add_text(slide, lead, 0.82, 3.18, 6.8, 1.0, 16, RGBColor(52,77,103))
    logo = Path(r"C:\Users\Administrator\Downloads\Makji_logo_img.png")
    if logo.exists():
        slide.shapes.add_picture(str(logo), Inches(8.55), Inches(0.95), width=Inches(3.8), height=Inches(3.8))
    else:
        mark = add_rect(slide, 8.75, 1.15, 3.35, 3.35, INK, INK)
        set_shape_text(mark, "MAKJI\nSTOCK", 25, WHITE, True, 0.2, MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER)
    meta = [clean_text(x) for x in section.xpath(".//*[contains(@class,'cover-meta')]//span")]
    if meta:
        add_text(slide, "   |   ".join(meta), 0.82, 5.35, 7.2, 0.35, 10, MUTED)
    add_text(slide, f"{page:02d}", 12.08, 7.06, 0.5, 0.18, 9, MUTED, False, PP_ALIGN.RIGHT)


def content_slide(prs, section, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
    sec = clean_text((section.xpath("./div[contains(@class,'top')]/div[contains(@class,'section')]") or [None])[0])
    eye = clean_text((section.xpath("./div[contains(@class,'eyebrow')]") or [None])[0])
    title = clean_text((section.xpath("./h2") or [None])[0])
    add_header(slide, sec, eye, title, page)

    formula = section.xpath(".//*[contains(concat(' ',normalize-space(@class),' '),' formula ')]")
    y0 = 2.25
    if formula:
        ftxt = clean_text(formula[0])
        fshape = add_rect(slide, 0.78, y0, 11.78, 0.85, NAVY, NAVY)
        set_shape_text(fshape, ftxt, 15, WHITE, True, 0.2, MSO_ANCHOR.MIDDLE)
        y0 = 3.28

    blocks = node_blocks(section)
    if not blocks:
        # Preserve any remaining meaningful text in an editable box.
        texts = [clean_text(e) for e in section.xpath("./p|./ul/li") if clean_text(e)]
        add_rich_block(slide, "핵심 내용", "\n".join(texts), 0.78, y0, 11.78, 3.85, True, False)
    else:
        n = len(blocks)
        cols = 2 if n <= 6 else 3
        rows = (n + cols - 1) // cols
        gap = 0.16
        total_h = 4.52 if y0 < 3 else 3.62
        bw = (11.78 - gap * (cols - 1)) / cols
        bh = (total_h - gap * (rows - 1)) / rows
        for i, (bt, body, accent, strong) in enumerate(blocks):
            col, row = i % cols, i // cols
            add_rich_block(slide, bt, body, 0.78 + col * (bw + gap), y0 + row * (bh + gap), bw, bh, accent, strong)

    source = section.xpath(".//*[contains(@class,'source')]")
    if source:
        add_text(slide, clean_text(source[0]), 0.78, 6.82, 10.9, 0.20, 7.5, MUTED)


def main():
    doc = html.fromstring(SOURCE.read_text(encoding="utf-8"))
    sections = doc.xpath("//section[contains(concat(' ',normalize-space(@class),' '),' slide ')]")
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    for page, section in enumerate(sections, 1):
        if "cover" in section.get("class", "").split():
            cover_slide(prs, section, page)
        else:
            content_slide(prs, section, page)
    prs.core_properties.title = "MAKJI STOCK 기업 대표 서비스 기획 발표 v0.1"
    prs.core_properties.subject = "HTML 원본 기반 수정 가능한 PowerPoint"
    prs.core_properties.author = "MAKJI"
    prs.save(OUTPUT)
    print(f"created: {OUTPUT}")
    print(f"slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
