from pptx import Presentation
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def extract_text_from_shape(shape):
    text = ""
    if hasattr(shape, "text") and shape.text.strip():
        text += shape.text.strip() + "\n"
    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                text += cell.text_frame.text.strip() + " "
            text += "\n"
    if shape.shape_type == 6:  # Group shape
        for child in shape.shapes:
            text += extract_text_from_shape(child)
    return text

def extract_text(file_path):
    try:
        prs = Presentation(file_path)
        with open("C:\\Users\\Administrator\\workspace\\makji\\extracted_text.txt", "w", encoding="utf-8") as f:
            for i, slide in enumerate(prs.slides):
                slide_text = ""
                for shape in slide.shapes:
                    slide_text += extract_text_from_shape(shape)
                if slide_text.strip():
                    f.write(f"--- Slide {i+1} ---\n{slide_text.strip()}\n\n")
                else:
                    f.write(f"--- Slide {i+1} ---\n[No Text Found]\n\n")
    except Exception as e:
        print(f"Error: {e}")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        extract_text(sys.argv[1])
